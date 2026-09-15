"""
UltraScrape - Exportador Microsoft Excel (.xlsx)
"""
import os
from typing import List
from core.models import ScrapedItem
from .base import BaseExporter

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class ExcelExporter(BaseExporter):
    def export(self, items: List[ScrapedItem]) -> bool:
        if not items:
            return False

        os.makedirs(os.path.dirname(os.path.abspath(self.output_path)), exist_ok=True)
        data = self.items_to_dict_list(items)

        if not OPENPYXL_AVAILABLE:
            # Fallback a CSV si openpyxl no está instalado
            csv_path = os.path.splitext(self.output_path)[0] + ".csv"
            from .csv_exporter import CsvExporter
            return CsvExporter(csv_path).export(items)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "UltraScrape_Export"

        # Headers
        headers = []
        for row in data:
            for k in row.keys():
                if k not in headers:
                    headers.append(k)

        ws.append(headers)

        # Formato de cabecera elegante (Cyber/Clean)
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Filas de datos
        for row in data:
            row_vals = [row.get(h, "") for h in headers]
            ws.append(row_vals)

        # Ajuste automático del ancho de columnas
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

        wb.save(self.output_path)
        return True
